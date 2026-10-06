from pathlib import Path
import pytest
from openpyxl.styles import Color, PatternFill, GradientFill
from openpyxl.writer.theme import theme_xml
from dienstplan_converter.excel_colors import normalizeFill, readTheme
from dienstplan_converter.excel_reader import readSchedule


@pytest.mark.parametrize(('color', 'expected'), [
    (Color(rgb='00CC66FF'), '#CC66FF'), (Color(rgb='FF123456'), '#123456'),
    (Color(theme=0), '#FFFFFF'), (Color(theme=1), '#000000'),
    (Color(theme=0, tint=-0.1499984740745262), '#D9D9D9'),
    (Color(theme=4, tint=0.7999816888943144), '#DCE6F2'),
    (Color(theme=9, tint=0.7999816888943144), '#FDEADA'),
    (Color(indexed=2), '#FF0000'), (Color(indexed=64), None),
    (Color(auto=True), None), (Color(theme=30), None),
    (Color(rgb='FF123456', tint=1), '#FFFFFF'),
    (Color(rgb='FF123456', tint=-1), '#000000'),
])
def testColorNormalization(color: Color, expected: str | None) -> None:
    assert normalizeFill(PatternFill('solid', fgColor=color), readTheme(theme_xml.encode())) == expected


def testNoFillAndUnsupportedColors() -> None:
    palette = readTheme(theme_xml.encode())
    assert normalizeFill(None, palette) is None
    assert normalizeFill(PatternFill(), palette) is None
    assert normalizeFill(PatternFill('darkGrid', fgColor='FF0000'), palette) is None
    assert normalizeFill(GradientFill(), palette) is None
    assert normalizeFill(PatternFill('solid', fgColor=Color(theme=0)), ()) is None
    assert normalizeFill(PatternFill('solid', fgColor=Color(indexed=0)), (), ('00123456',)) == '#123456'


@pytest.mark.parametrize(('month', 'day', 'expected'), [
    (1, 1, ('#FF7C80', '#FF7C80', '#FF7C80', '#FF7C80')),
    (1, 2, ('#FFFFFF', '#FFFFFF', '#DEEBF7', '#DEEBF7')),
    (4, 2, ('#CC66FF', '#92D050', '#FFFFFF', '#FFFFFF')),
    (8, 31, ('#FFFFFF', '#FFFFFF', '#E2F0D9', '#E2F0D9')),
    (9, 1, ('#FFFFFF', '#FFFFFF', '#92D050', '#92D050')),
    (10, 19, ('#FFFFFF', '#FFFFFF', '#FFFF00', '#FFFFFF')),
])
def testFixtureColors(month: int, day: int, expected: tuple[str, ...]) -> None:
    path = Path(__file__).parent / 'fixtures' / f'dienstplan_2026_{month:02}.xlsx'
    schedule = readSchedule(path)
    entry = schedule.days[day - 1]
    assert (entry.headerColor, entry.weekdayColor, entry.primaryColor, entry.secondaryColor) == expected
    assert all(color is not None for entry in schedule.days for color in
               (entry.headerColor, entry.weekdayColor, entry.primaryColor, entry.secondaryColor))


def testThemeMappingUsesWorkbookValues() -> None:
    xml = theme_xml.replace('5B9BD5', '123456').encode()
    # Default openpyxl theme has accent1=4F81BD, fixtures have accent1=5B9BD5.
    xml = xml.replace(b'4F81BD', b'123456')
    palette = readTheme(xml)
    assert normalizeFill(PatternFill('solid', fgColor=Color(theme=4)), palette) == '#123456'
