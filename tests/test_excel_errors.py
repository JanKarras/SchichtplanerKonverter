from pathlib import Path
import pytest
from openpyxl import load_workbook
from dienstplan_converter.errors import ConversionError
from dienstplan_converter.excel_reader import readSchedule


@pytest.mark.parametrize(('cell', 'value', 'message'), [
    ('N1', None, 'Monat'), ('N1', 'April', 'Jahr'),
    ('B4', None, 'Mitarbeiter/Tag'), ('B7', 'Müller', 'Klein'),
    ('D3', 9, 'Tagesfolge'), ('AF3', None, 'Tagesfolge'),
    ('B8', 'Müller', 'zweite Dienstplanzeile'),
])
def testInvalidStructure(sampleFile: Path, cell: str, value: object, message: str) -> None:
    workbook = load_workbook(sampleFile)
    workbook.active[cell] = value
    workbook.save(sampleFile)
    workbook.close()
    with pytest.raises(ConversionError, match=message):
        readSchedule(sampleFile)


@pytest.mark.parametrize('filename', ['missing.xlsx', 'wrong.xls', 'broken.xlsx'])
def testInvalidFiles(tmp_path: Path, filename: str) -> None:
    path = tmp_path / filename
    if filename == 'broken.xlsx':
        path.write_text('invalid')
    with pytest.raises(ConversionError):
        readSchedule(path)


def testSecondSheet(sampleFile: Path) -> None:
    workbook = load_workbook(sampleFile)
    workbook.create_sheet('Hinweise', 0)['A1'] = 'Notizen'
    workbook.save(sampleFile)
    workbook.close()
    assert readSchedule(sampleFile).employee == 'Klein'


def testAmbiguousWorkbook(sampleFile: Path) -> None:
    workbook = load_workbook(sampleFile)
    workbook.copy_worksheet(workbook.active)
    workbook.save(sampleFile)
    workbook.close()
    with pytest.raises(ConversionError, match='Mehrere'):
        readSchedule(sampleFile)
