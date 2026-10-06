from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from dienstplan_converter.excel_reader import readSchedule


def testBlankPositionsAndNoFill(sampleFile: Path) -> None:
    workbook = load_workbook(sampleFile)
    sheet = workbook.active
    sheet['C7'].fill = PatternFill('solid', fgColor='FF123456')
    sheet['D7'].fill = PatternFill('solid', fgColor='FF654321')
    sheet['E8'].fill = PatternFill('solid', fgColor='FFABCDEF')
    workbook.save(sampleFile)
    workbook.close()
    days = readSchedule(sampleFile).days
    assert days[0].primaryColor == '#123456'
    assert days[1].primaryValue is None
    assert days[1].primaryColor == '#654321'
    assert days[2].secondaryColor == '#ABCDEF'
    assert days[2].primaryColor is None
    assert days[3].headerColor is None
    assert days[3].weekdayColor is None
    assert days[3].secondaryColor is None


def testFollowingMonthAndExtraColors(sampleFile: Path) -> None:
    workbook = load_workbook(sampleFile)
    sheet = workbook.active
    sheet['N1'] = 'Januar 2026'
    sheet['AG3'] = 31
    sheet['AH3'] = 1
    for row in (3, 4, 7, 8):
        sheet.cell(row, 34).fill = PatternFill('solid', fgColor='FF112233')
        sheet.cell(row, 37).fill = PatternFill('solid', fgColor='FF445566')
    workbook.save(sampleFile)
    workbook.close()
    days = readSchedule(sampleFile).days
    assert len(days) == 31
    assert not any(color in ('#112233', '#445566') for day in days for color in
                   (day.headerColor, day.weekdayColor, day.primaryColor, day.secondaryColor))


def testShiftedColors(sampleFile: Path) -> None:
    workbook = load_workbook(sampleFile)
    sheet = workbook.active
    sheet['C3'].fill = PatternFill('solid', fgColor='FF123456')
    sheet['C8'].fill = PatternFill('solid', fgColor='FF654321')
    sheet.insert_rows(2)
    sheet.insert_cols(1)
    workbook.save(sampleFile)
    workbook.close()
    first = readSchedule(sampleFile).days[0]
    assert first.headerColor == '#123456'
    assert first.secondaryColor == '#654321'
