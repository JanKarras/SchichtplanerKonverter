from pathlib import Path
from unittest.mock import Mock
import pytest
from openpyxl import load_workbook
from dienstplan_converter.excel_reader import readSchedule
from dienstplan_converter import excel_reader


def testExplicitFollowingMonth(sampleFile: Path) -> None:
    workbook = load_workbook(sampleFile)
    sheet = workbook.active
    sheet['N1'] = 'Januar 2026'
    sheet['AG3'] = 31
    sheet['AH3'] = 1
    sheet['AG8'] = 'LAST'
    sheet['AH7'] = 'FEBRUARY'
    sheet['AH8'] = 'FEBRUARY'
    workbook.save(sampleFile)
    workbook.close()
    schedule = readSchedule(sampleFile)
    assert len(schedule.days) == 31
    assert schedule.days[-1].secondaryValue == 'LAST'
    assert not any(day.primaryValue == 'FEBRUARY' or day.secondaryValue == 'FEBRUARY' for day in schedule.days)


def testStrictReadBounds(monkeypatch, sampleFile: Path) -> None:
    realWorkbook = load_workbook(sampleFile, read_only=True, data_only=True)
    rows = tuple(realWorkbook.active.iter_rows(min_row=1, max_row=30, min_col=1, max_col=80, values_only=False))
    realWorkbook.close()
    sheet = Mock()
    sheet.iter_rows.return_value = iter(rows)
    workbook = Mock(worksheets=[sheet])
    monkeypatch.setattr(excel_reader, 'load_workbook', lambda *args, **kwargs: workbook)
    assert readSchedule(sampleFile).month == 4
    sheet.iter_rows.assert_called_once_with(min_row=1, max_row=30, min_col=1, max_col=80, values_only=False)
    workbook.close.assert_called_once()


@pytest.mark.parametrize(('year', 'count'), [(2026, 28), (2028, 29)])
def testFebruaryParser(sampleFile: Path, year: int, count: int) -> None:
    workbook = load_workbook(sampleFile)
    sheet = workbook.active
    sheet['N1'] = f'Februar {year}'
    workbook.save(sampleFile)
    workbook.close()
    assert len(readSchedule(sampleFile).days) == count


def testStableRealOctoberValues() -> None:
    schedule = readSchedule(Path(__file__).parent / 'fixtures' / 'dienstplan_2026_10.xlsx')
    assert schedule.days[6].primaryValue == '/F'
    assert schedule.days[6].secondaryValue == '/S'
    assert schedule.days[25].primaryValue == '07-15:30'
    assert schedule.days[25].secondaryValue == 'ASH'
