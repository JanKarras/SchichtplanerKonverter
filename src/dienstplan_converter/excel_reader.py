import datetime
from pathlib import Path
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.exceptions import InvalidFileException

from .calendar_utils import WEEKDAY_NAMES
from .excel_colors import normalizeFill, readTheme
from .errors import ConversionError
from .excel_structure import SEARCH_COLUMNS, SEARCH_ROWS, findMonth, findStructure
from .models import DayEntry, MonthSchedule


def cellText(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.time)):
        return value.strftime('%H:%M')
    return str(value)


def parseSheet(cellRows: tuple[tuple[ReadOnlyCell | EmptyCell, ...], ...],
               theme: tuple[str | None, ...], indexed: tuple[str, ...]) -> MonthSchedule:
    rows = tuple(tuple(cell.value for cell in row) for row in cellRows)
    month, year = findMonth(rows)
    employeeRow, startColumn, dayCount, dayRow, weekdayRow = findStructure(rows, month, year)
    days = tuple(DayEntry(
        day=day,
        weekday=WEEKDAY_NAMES[datetime.date(year, month, day).weekday()],
        primaryValue=cellText(rows[employeeRow][startColumn + day - 1]),
        secondaryValue=cellText(rows[employeeRow + 1][startColumn + day - 1]),
        headerColor=normalizeFill(cellRows[dayRow][startColumn + day - 1].fill, theme, indexed),
        weekdayColor=normalizeFill(cellRows[weekdayRow][startColumn + day - 1].fill, theme, indexed),
        primaryColor=normalizeFill(cellRows[employeeRow][startColumn + day - 1].fill, theme, indexed),
        secondaryColor=normalizeFill(cellRows[employeeRow + 1][startColumn + day - 1].fill, theme, indexed),
    ) for day in range(1, dayCount + 1))
    return MonthSchedule('Klein', year, month, days)


def readSchedule(inputPath: Path) -> MonthSchedule:
    if inputPath.suffix.lower() != '.xlsx':
        raise ConversionError('Bitte wählen Sie eine Excel-Datei mit der Endung .xlsx.')
    if not inputPath.is_file():
        raise ConversionError('Die Excel-Datei existiert nicht oder ist nicht lesbar.')
    try:
        workbook = load_workbook(inputPath, read_only=True, data_only=True, keep_links=False)
        try:
            theme = readTheme(workbook.loaded_theme)
            # openpyxl 3.x stores the workbook-specific indexed palette here.
            indexed = tuple(workbook._colors) if isinstance(workbook._colors, (list, tuple)) else COLOR_INDEX
            schedules = []
            failures = []
            for sheet in workbook.worksheets:
                rows = tuple(sheet.iter_rows(min_row=1, max_row=SEARCH_ROWS,
                                            min_col=1, max_col=SEARCH_COLUMNS, values_only=False))
                try:
                    schedules.append(parseSheet(rows, theme, indexed))
                except ConversionError as error:
                    failures.append(error)
            if len(schedules) == 1:
                return schedules[0]
            if len(schedules) > 1:
                raise ConversionError('Mehrere Dienstpläne gefunden. Bitte verwenden Sie eine Datei mit genau einem Monatsplan.')
            raise failures[0] if failures else ConversionError('Kein lesbares Arbeitsblatt gefunden.')
        finally:
            workbook.close()
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError, EOFError) as error:
        raise ConversionError('Die Excel-Datei konnte nicht gelesen werden. Sie ist möglicherweise beschädigt oder keine gültige .xlsx-Datei.') from error
